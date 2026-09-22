"""XLSX -> CSV conversion: a standalone pre-processing step, never part of
the ingest/adapter pipeline itself. `ingest.py`'s own CSV-reading code
(`open_csv`/`iter_csv_rows`) needs NO changes for this module to exist, and
never will -- an `.xlsx` source is converted to real CSV file(s) FIRST, via
`rhino xlsx-to-csv`, and everything after that point is the exact same
native/defender/bluepeak/configured-adapter pipeline every other CSV source
already goes through, completely unchanged.

Why conversion, not a native XLSX-reading adapter (docs/handoff.md's own
"XLSX intake" survey named four real, unresolved technical constraints and
deliberately decided nothing -- CLAUDE.md's dated entry on this module
records the actual decision): typed cells (openpyxl returns real
int/float/date/bool objects, not strings, and every parser downstream of a
CSV row assumes `str` throughout the ingest layer), sheet selection
(`Source.layout` has no dimension for "which sheet"), ragged-row protection
(CSV's own line-oriented ragged-row detection has no XLSX equivalent), and
a genuine tension in openpyxl's own API between memory-bounded streaming
and merged-cell detection. Every one of those would have had to be solved
INSIDE the shared ingest/probe/configured-adapter layer every other source
format also depends on, for a capability only some sources need. Converting
first solves each problem HERE, once, in isolation:

- Typed cells: resolved by `_stringify` below, once, rather than by making
  every downstream parser (`ParsedMapping`, `_apply_case`, the delimiter/
  row-shape logic) newly tolerant of non-`str` input.
- Sheet selection: a CLI concern (`--sheet`) of THIS tool, never a question
  the ingest schema (`Source.layout`) has to grow a dimension for.
- Ragged-row protection: not reimplemented at all -- the OUTPUT of this
  module is a real CSV file, so `ingest.py`'s existing, already-tested
  ragged-row detection (`iter_csv_rows`) applies to it for free, downstream,
  exactly as it does to any hand-written CSV.
- Memory vs. merge-detection: NOT resolved by picking one and giving up the
  other -- this module always loads a full worksheet (never openpyxl's
  `read_only=True` streaming mode), because CLAUDE.md's own stated target
  scale (hundreds to thousands of findings) is trivially small for a full
  in-memory worksheet load; the tension the survey named only bites at a
  scale this project isn't targeting. Documented here as a real, named scope
  limit (CLAUDE.md's own "no silent caps" discipline), not silently assumed
  away: a workbook with genuinely hundreds of thousands of rows would need a
  different tool.

Two things this module refuses rather than silently mis-converts, matching
the project's not-collected/refuse-rather-than-guess discipline applied to
a new kind of input: a MERGED cell region (a flat CSV grid cannot losslessly
represent one -- every cell but the top-left reads back empty, which is
real, silent data loss unless a human explicitly opts into it via
`unmerge=True`), and a FORMULA cell with no cached calculated value (openpyxl
with `data_only=True` returns `None` for this -- indistinguishable from a
genuinely blank cell without also checking the formula-preserving load,
which this module always does, precisely to tell the two apart and refuse
the latter with a real, actionable message instead of writing a silently
blank CSV cell where a real value belongs).
"""

from __future__ import annotations

import csv
import datetime
from dataclasses import dataclass
from pathlib import Path
from typing import Any


class XlsxConversionError(Exception):
    """Raised for anything this module refuses to guess about -- a merged
    region, an uncalculated formula, an unreadable workbook, or an output
    directory this module won't silently overwrite. Never raised for a
    genuinely empty cell or sheet; those convert to blank CSV cells/empty
    files, the honest representation of "nothing was there."""


@dataclass(frozen=True)
class ConvertedSheet:
    """One sheet's conversion result -- returned even though the CSV is
    also written to disk, so a caller (the CLI) can report row/column
    counts without re-reading the file it just wrote."""

    sheet_name: str
    csv_path: Path
    row_count: int
    column_count: int


def _stringify(value: Any) -> str:
    """The one place a typed openpyxl cell value becomes the string a CSV
    cell, and everything downstream of it, expects. Deliberately narrow --
    this is not a general Python-to-string formatter, it exists to match
    exactly the conventions the rest of this project's CSV fixtures already
    use (confirmed against data/demo/assets.csv: `True`/`False`, plain
    digit strings, ISO dates), so a converted file reads identically to a
    hand-written one.

    - None (a genuinely blank cell) -> "" -- the same blank a hand-authored
      CSV would have, never a guessed value.
    - bool -> "True"/"False" (Python's own str(bool), which already matches
      the fixture convention -- checked BEFORE the int/float branch below,
      since bool is a subclass of int in Python and would otherwise be
      caught by it first, producing "1"/"0" instead).
    - A float that represents a whole number (5.0) -> "5", not "5.0" --
      openpyxl/Excel's own round-trip already normalizes most whole-number
      floats to int on save (confirmed live), but a value arriving as a
      genuine float here (e.g. from a formula's cached result) still must
      not produce a string a strict int-typed schema field would reject.
    - int / other float -> str(value).
    - datetime.datetime at exactly midnight, or datetime.date -> plain ISO
      date (YYYY-MM-DD), matching detected_date's own convention in every
      existing fixture. A datetime with a real time component keeps it
      (full ISO 8601) rather than silently discarding it.
    - Anything else (str, or a type openpyxl doesn't otherwise produce) ->
      str(value), unchanged -- never stripped or otherwise rewritten, so
      this never alters a value a human actually typed.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, datetime.datetime):
        return value.date().isoformat() if value.time() == datetime.time.min else value.isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value)


def _check_no_uncalculated_formulas(ws_values: Any, ws_formulas: Any, sheet_name: str) -> None:
    """A formula cell openpyxl can't report a cached value for reads as
    `None` under `data_only=True` -- identical, from that load alone, to a
    genuinely blank cell. Told apart here by cross-referencing the SAME
    cell's `data_type` in the formula-preserving load ("f" means "this is a
    formula"); refused by name and location rather than silently written as
    an empty CSV cell, which would be real, silent data loss dressed up as
    "no data collected" -- a different and false claim."""
    problems: list[str] = []
    for row_v, row_f in zip(ws_values.iter_rows(), ws_formulas.iter_rows()):
        for cell_v, cell_f in zip(row_v, row_f):
            if cell_v.value is None and cell_f.data_type == "f":
                problems.append(f"{sheet_name}!{cell_f.coordinate} (formula {cell_f.value!r})")
    if problems:
        raise XlsxConversionError(
            f"{len(problems)} cell(s) have a formula with no cached calculated value, so the real "
            f"value to convert is unknown, not blank: {', '.join(problems)}. Open the workbook in "
            "Excel (or LibreOffice) and save it once so a value gets cached, or replace the formula "
            "with its literal value, then convert again."
        )


def _check_no_merged_cells(ws: Any, sheet_name: str, *, unmerge: bool) -> None:
    """A merged region has exactly one real value (the top-left cell) and
    every other cell in the range reads back empty -- flattening that into
    a CSV grid without saying so would silently blank out cells that were
    never actually empty. Refused by default, naming every merged range;
    `unmerge=True` is the explicit, human-chosen opt-in to accept that
    lossy-but-sometimes-wanted flattening (each cell in the range keeps
    reading empty except the top-left -- this function does not fill the
    range, it only stops REFUSING the conversion when a human has said
    that's fine)."""
    if unmerge or not ws.merged_cells.ranges:
        return
    ranges = sorted(str(r) for r in ws.merged_cells.ranges)
    raise XlsxConversionError(
        f"{sheet_name!r} has {len(ranges)} merged cell region(s): {', '.join(ranges)}. A flat CSV grid "
        "cannot represent a merge losslessly -- every cell but the top-left of each region would read "
        "back empty, which looks like missing data, not a merge. Unmerge the cells in the source "
        "workbook and re-save, or pass unmerge=True (--allow-merged-cells on the CLI) to accept that "
        "every cell but each region's top-left converts to a blank."
    )


def _sheet_names(workbook: Any, requested: list[str] | None) -> list[str]:
    if requested is None:
        return list(workbook.sheetnames)
    missing = [name for name in requested if name not in workbook.sheetnames]
    if missing:
        raise XlsxConversionError(
            f"sheet(s) {missing} not found in this workbook. Real sheets: {list(workbook.sheetnames)}"
        )
    return requested


def convert_workbook(
    xlsx_path: Path,
    out_dir: Path,
    *,
    sheets: list[str] | None = None,
    overwrite: bool = False,
    unmerge: bool = False,
) -> list[ConvertedSheet]:
    """Converts every sheet in `xlsx_path` (or only `sheets`, if given) to
    a same-named CSV file inside `out_dir` -- `<out_dir>/<sheet name>.csv`
    per sheet, ready to be pointed at as a `--data` directory (native
    format, if the sheets are already named/shaped like `assets`/
    `findings`) or as `--assets-file`/`--findings-file`/
    `rhino adapt propose`'s own input, unchanged from any other CSV source
    from this point on.

    `out_dir` is refused if it already exists and is non-empty, unless
    `overwrite=True` -- the same "don't silently clobber what's there"
    caution this project applies to every other write path. Created if it
    doesn't exist yet.

    Loads the workbook TWICE, deliberately: once with `data_only=True` (the
    values a human sees), once without (to tell a formula's uncalculated
    `None` apart from a genuinely blank cell -- see
    `_check_no_uncalculated_formulas`). Never uses openpyxl's `read_only=
    True` streaming mode -- see this module's own docstring for why that's
    a real, named scope limit rather than an oversight."""
    import openpyxl

    if not xlsx_path.is_file():
        raise XlsxConversionError(f"{xlsx_path}: no such file")

    try:
        wb_values = openpyxl.load_workbook(xlsx_path, data_only=True)
    except Exception as exc:  # openpyxl raises several distinct exception types for a bad/corrupt file
        raise XlsxConversionError(f"{xlsx_path}: could not open as an XLSX workbook: {exc}") from exc
    wb_formulas = openpyxl.load_workbook(xlsx_path, data_only=False)

    names = _sheet_names(wb_values, sheets)

    if out_dir.exists():
        if not out_dir.is_dir():
            raise XlsxConversionError(f"{out_dir}: exists and is not a directory")
        existing = list(out_dir.iterdir())
        if existing and not overwrite:
            raise XlsxConversionError(
                f"{out_dir}: already exists and is not empty ({len(existing)} entr"
                f"{'y' if len(existing) == 1 else 'ies'} on file) -- pass overwrite=True "
                "(--overwrite on the CLI) to convert into it anyway."
            )
    else:
        out_dir.mkdir(parents=True)

    results: list[ConvertedSheet] = []
    for name in names:
        ws_values = wb_values[name]
        ws_formulas = wb_formulas[name]
        _check_no_merged_cells(ws_values, name, unmerge=unmerge)
        _check_no_uncalculated_formulas(ws_values, ws_formulas, name)

        rows = [[_stringify(cell.value) for cell in row] for row in ws_values.iter_rows()]
        # Trailing wholly-empty rows/columns are common in a real export
        # (Excel's own "used range" tracking is generous) and would
        # otherwise become a ragged-looking block of blank CSV cells or
        # blank trailing rows -- trimmed here, once, rather than relied on
        # to look harmless everywhere downstream.
        while rows and all(cell == "" for cell in rows[-1]):
            rows.pop()
        col_count = max((len(r) for r in rows), default=0)
        while col_count and all((r[col_count - 1] if col_count - 1 < len(r) else "") == "" for r in rows):
            col_count -= 1
        rows = [r[:col_count] for r in rows]

        csv_path = out_dir / f"{name}.csv"
        with csv_path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerows(rows)

        results.append(ConvertedSheet(sheet_name=name, csv_path=csv_path, row_count=len(rows), column_count=col_count))

    return results
