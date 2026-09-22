"""Coverage for the `rhino xlsx-to-csv` command's own wiring -- argument
parsing, exit codes, and what gets printed. The conversion logic itself is
`xlsx_convert.py`'s own concern, fully covered in `tests/test_xlsx_convert
.py`; these tests exist to confirm the CLI reaches it correctly and reports
what happened, not to re-prove the conversion rules.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from rhinosecure.cli import main


@pytest.fixture
def simple_xlsx(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "assets"
    ws.append(["asset_id", "hostname"])
    ws.append(["A01", "DC01"])
    path = tmp_path / "source.xlsx"
    wb.save(path)
    return path


def test_xlsx_to_csv_converts_and_reports_each_sheet(simple_xlsx: Path, tmp_path: Path, capsys):
    out_dir = tmp_path / "out"
    exit_code = main(["xlsx-to-csv", str(simple_xlsx), "--out-dir", str(out_dir)])
    assert exit_code == 0
    out = capsys.readouterr().out
    assert "assets: 2 row(s), 2 column(s)" in out
    assert str(out_dir / "assets.csv") in out
    assert "1 sheet(s) converted" in out
    assert (out_dir / "assets.csv").exists()


def test_xlsx_to_csv_sheet_flag_is_repeatable(tmp_path: Path, capsys):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "one"
    ws1.append(["a"])
    ws2 = wb.create_sheet("two")
    ws2.append(["b"])
    ws3 = wb.create_sheet("three")
    ws3.append(["c"])
    xlsx_path = tmp_path / "source.xlsx"
    wb.save(xlsx_path)

    out_dir = tmp_path / "out"
    exit_code = main([
        "xlsx-to-csv", str(xlsx_path), "--out-dir", str(out_dir), "--sheet", "one", "--sheet", "three",
    ])
    assert exit_code == 0
    assert (out_dir / "one.csv").exists()
    assert (out_dir / "three.csv").exists()
    assert not (out_dir / "two.csv").exists()


def test_xlsx_to_csv_a_conversion_error_exits_1_with_a_clean_message_not_a_traceback(simple_xlsx: Path, tmp_path: Path, capsys):
    exit_code = main([
        "xlsx-to-csv", str(simple_xlsx), "--out-dir", str(tmp_path), "--sheet", "does-not-exist",
    ])
    assert exit_code == 1
    err = capsys.readouterr().err
    assert "xlsx-to-csv:" in err
    assert "does-not-exist" in err


def test_xlsx_to_csv_missing_source_file_exits_1(tmp_path: Path, capsys):
    exit_code = main([
        "xlsx-to-csv", str(tmp_path / "nope.xlsx"), "--out-dir", str(tmp_path / "out"),
    ])
    assert exit_code == 1
    assert "no such file" in capsys.readouterr().err


def test_xlsx_to_csv_merged_cells_refused_by_default_via_the_cli(tmp_path: Path, capsys):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "assets"
    ws["A1"] = "header"
    ws.merge_cells("A1:B1")
    xlsx_path = tmp_path / "source.xlsx"
    wb.save(xlsx_path)

    exit_code = main(["xlsx-to-csv", str(xlsx_path), "--out-dir", str(tmp_path / "out")])
    assert exit_code == 1
    assert "merged" in capsys.readouterr().err.lower()


def test_xlsx_to_csv_allow_merged_cells_flag_lets_it_through(tmp_path: Path, capsys):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "assets"
    ws["A1"] = "header"
    ws.merge_cells("A1:B1")
    xlsx_path = tmp_path / "source.xlsx"
    wb.save(xlsx_path)

    exit_code = main([
        "xlsx-to-csv", str(xlsx_path), "--out-dir", str(tmp_path / "out"), "--allow-merged-cells",
    ])
    assert exit_code == 0


def test_xlsx_to_csv_out_dir_required(simple_xlsx: Path):
    with pytest.raises(SystemExit):
        main(["xlsx-to-csv", str(simple_xlsx)])


def test_xlsx_to_csv_overwrite_flag_reaches_the_conversion(simple_xlsx: Path, tmp_path: Path):
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    (out_dir / "unrelated.txt").write_text("x", encoding="utf-8")

    refused = main(["xlsx-to-csv", str(simple_xlsx), "--out-dir", str(out_dir)])
    assert refused == 1

    allowed = main(["xlsx-to-csv", str(simple_xlsx), "--out-dir", str(out_dir), "--overwrite"])
    assert allowed == 0
